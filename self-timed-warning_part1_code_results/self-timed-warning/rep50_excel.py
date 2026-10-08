"""Excel report for the repeated 50/168 test (supervisor's A-B-C-D-E layout) -> out/rep50/Warning_model_50_engine_tests.xlsx"""
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
R = pd.read_csv('out/rep50/rep50_results.csv'); NR = R.repetition.nunique()
F = 'Arial'; HF = PatternFill('solid', fgColor='1F4E79'); HFont = Font(name=F, bold=True, color='FFFFFF'); B = Font(name=F); BB = Font(name=F, bold=True)
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
GREEN = PatternFill('solid', fgColor='E2F0D9'); RED = PatternFill('solid', fgColor='FBE3E1'); AMBER = PatternFill('solid', fgColor='FFF2CC')
def head(ws, r, vals, widths=None):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = HFont; c.fill = HF; c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    if widths:
        for j, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(j)].width = w
wb = Workbook()

# ---------- All engines (long table), E and result as formulas
al = wb.active; al.title = 'All tests'
head(al, 1, ['Repetition', 'Engine', 'C: actual life (cycles)', 'Warning issued at cycle', 'Predicted cycles left at warning',
             'D: predicted life (cycles)', 'E = D - C (cycles)', 'Result'], [11, 9, 13, 13, 14, 14, 13, 12]); al.row_dimensions[1].height = 45
for i, x in enumerate(R.sort_values(['repetition', 'engine']).itertuples(), 2):
    al.cell(i, 1, int(x.repetition)); al.cell(i, 2, int(x.engine)); al.cell(i, 3, int(x.actual_life_C))
    if pd.isna(x.predicted_life_D):
        al.cell(i, 4, 'none'); al.cell(i, 5, '-'); al.cell(i, 6, '-'); al.cell(i, 7, '-'); al.cell(i, 8, 'MISSED')
    else:
        al.cell(i, 4, int(x.warning_cycle)); al.cell(i, 5, float(x.predicted_cycles_left)); al.cell(i, 5).number_format = '0.0'; al.cell(i, 6, f'=D{i}+E{i}')
        al.cell(i, 7, f'=F{i}-C{i}'); al.cell(i, 8, f'=IF(G{i}>2,"LATE",IF(G{i}<-13,"EARLY","IN"))')
    for j in range(1, 9):
        c = al.cell(i, j); c.font = B; c.border = BOX; c.alignment = Alignment(horizontal='center')
    al.cell(i, 6).number_format = '0.0'; al.cell(i, 7).number_format = '0.0'
n_all = len(R) + 1; al.freeze_panes = 'A2'
al.conditional_formatting.add(f'H2:H{n_all}', CellIsRule(operator='equal', formula=['"IN"'], fill=GREEN))
al.conditional_formatting.add(f'H2:H{n_all}', CellIsRule(operator='notEqual', formula=['"IN"'], fill=RED))

# ---------- Summary
s = wb.create_sheet('Summary', 0)
for col, w in zip('ABCDEFG', [34, 14, 12, 12, 12, 14, 30]): s.column_dimensions[col].width = w
s['A1'] = 'Warning model, repeated random 50-engine tests (PHM08)'; s['A1'].font = Font(name=F, bold=True, size=14)
notes = ['A: 50 engines drawn at random as the test set; B: the other 168 engines train the model (settings tuned on the 168 only).',
         'C: actual life; D: predicted life = warning cycle + predicted cycles left; E = D - C. Correct when -13 <= E <= +2.',
         'The warning is issued when the model predicts about 20 cycles left (5 cycles in a row). Repeated 20 times with new random test sets.']
for k, t in enumerate(notes, 2): s.cell(k, 1, t).font = Font(name=F, italic=True, color='555555')
r0 = 6; head(s, r0, ['Repetition', 'In window (of 50)', 'Late', 'Early', 'Missed', '% in window', 'Meets 48-of-50?'])
for r in range(1, NR + 1):
    i = r0 + r; s.cell(i, 1, r)
    s.cell(i, 2, f"=COUNTIFS('All tests'!$A$2:$A${n_all},A{i},'All tests'!$H$2:$H${n_all},\"IN\")")
    s.cell(i, 3, f"=COUNTIFS('All tests'!$A$2:$A${n_all},A{i},'All tests'!$H$2:$H${n_all},\"LATE\")")
    s.cell(i, 4, f"=COUNTIFS('All tests'!$A$2:$A${n_all},A{i},'All tests'!$H$2:$H${n_all},\"EARLY\")")
    s.cell(i, 5, f"=COUNTIFS('All tests'!$A$2:$A${n_all},A{i},'All tests'!$H$2:$H${n_all},\"MISSED\")")
    s.cell(i, 6, f'=B{i}/50'); s.cell(i, 6).number_format = '0.0%'; s.cell(i, 7, f'=IF(B{i}>=48,"Yes","No")')
    for j in range(1, 8): s.cell(i, j).font = B; s.cell(i, j).border = BOX; s.cell(i, j).alignment = Alignment(horizontal='center')
last = r0 + NR
s.conditional_formatting.add(f'G{r0+1}:G{last}', CellIsRule(operator='equal', formula=['"Yes"'], fill=GREEN))
s.conditional_formatting.add(f'G{r0+1}:G{last}', CellIsRule(operator='equal', formula=['"No"'], fill=AMBER))
k = last + 2; head(s, k, ['Across all repetitions', 'Value'])
stats = [('Average engines in window (of 50)', f'=AVERAGE(B{r0+1}:B{last})', '0.0'),
         ('Average % in window', f'=AVERAGE(F{r0+1}:F{last})', '0.0%'),
         ('Standard deviation (engines)', f'=STDEV(B{r0+1}:B{last})', '0.0'),
         ('Worst repetition (engines in window)', f'=MIN(B{r0+1}:B{last})', '0'),
         ('Best repetition (engines in window)', f'=MAX(B{r0+1}:B{last})', '0'),
         ('Repetitions meeting 48 of 50', f'=COUNTIF(G{r0+1}:G{last},"Yes")', '0'),
         ('Repetitions with at least 46 of 50', f'=COUNTIF(B{r0+1}:B{last},">=46")', '0'),
         ('Average late + missed per 50', f'=AVERAGE(C{r0+1}:C{last})+AVERAGE(E{r0+1}:E{last})', '0.0'),
         ('Average too early per 50', f'=AVERAGE(D{r0+1}:D{last})', '0.0'),
         ('Median error E (all tests, cycles)', f"=MEDIAN('All tests'!G2:G{n_all})", '0.0')]
for j, (a, f, fmt) in enumerate(stats, k + 1):
    s.cell(j, 1, a).font = B; c = s.cell(j, 2, f); c.font = BB; c.number_format = fmt; s.cell(j, 1).border = c.border = BOX
s.cell(k + len(stats) + 2, 1, 'Each engine appears in several test sets across repetitions; repetitions are not independent.').font = Font(name=F, italic=True, color='555555')

# ---------- Repetition 1 in the supervisor's layout
ex = wb.create_sheet('Example repetition 1', 1)
ex['A1'] = 'Repetition 1: the 50 test engines (supervisor layout C, D, E)'; ex['A1'].font = Font(name=F, bold=True, size=13)
head(ex, 3, ['No.', 'Engine', 'C: actual life', 'D: predicted life', 'E = D - C', 'Result'], [7, 10, 14, 16, 12, 12])
R1 = R[R.repetition == 1].sort_values('engine')
for i, x in enumerate(R1.itertuples(), 4):
    ex.cell(i, 1, i - 3); ex.cell(i, 2, int(x.engine)); ex.cell(i, 3, int(x.actual_life_C))
    if pd.isna(x.predicted_life_D): ex.cell(i, 4, 'no warning'); ex.cell(i, 5, '-'); ex.cell(i, 6, 'MISSED')
    else:
        ex.cell(i, 4, float(x.predicted_life_D)); ex.cell(i, 4).number_format = '0.0'; ex.cell(i, 5, f'=D{i}-C{i}'); ex.cell(i, 5).number_format = '0.0'
        ex.cell(i, 6, f'=IF(E{i}>2,"LATE",IF(E{i}<-13,"EARLY","IN"))')
    for j in range(1, 7): ex.cell(i, j).font = B; ex.cell(i, j).border = BOX; ex.cell(i, j).alignment = Alignment(horizontal='center')
e1 = 3 + len(R1)
ex.conditional_formatting.add(f'F4:F{e1}', CellIsRule(operator='equal', formula=['"IN"'], fill=GREEN))
ex.conditional_formatting.add(f'F4:F{e1}', CellIsRule(operator='notEqual', formula=['"IN"'], fill=RED))
ex.cell(e1 + 2, 1, 'In window:').font = BB; ex.cell(e1 + 2, 3, f'=COUNTIF(F4:F{e1},"IN")').font = BB; ex.cell(e1 + 2, 4, 'of 50').font = B
wb.save('out/rep50/Warning_model_50_engine_tests.xlsx'); print('saved')
